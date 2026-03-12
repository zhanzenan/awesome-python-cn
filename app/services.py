from __future__ import annotations

import re
from datetime import date
from pathlib import Path
from typing import Dict, List, Optional

from pydantic import BaseModel, Field


class ExtractedMaterial(BaseModel):
    people: List[str] = Field(default_factory=list)
    facts: List[str] = Field(default_factory=list)
    evidences: List[str] = Field(default_factory=list)
    procedure_nodes: List[str] = Field(default_factory=list)
    ocr_confidence: float = 0.8
    raw_preview: str = ""


class GenerationResult(BaseModel):
    title: str
    sections: List[Dict[str, str]]
    references_used: List[str]
    law_suggestions: List[Dict]
    traceability: List[Dict[str, str]]


def _read_text(file_path: Path) -> str:
    suffix = file_path.suffix.lower()
    if suffix in {".txt", ".md"}:
        return file_path.read_text(encoding="utf-8", errors="ignore")
    if suffix in {".json"}:
        return file_path.read_text(encoding="utf-8", errors="ignore")
    return f"[二进制文件]{file_path.name}，需OCR识别。"


def extract_material(file_path: Path) -> ExtractedMaterial:
    text = _read_text(file_path)
    names = re.findall(r"[\u4e00-\u9fa5]{2,4}(?:某|甲|乙)?", text)
    times = re.findall(r"\d{4}年\d{1,2}月\d{1,2}日", text)
    facts = [f"案发时间涉及：{t}" for t in times[:3]]
    evidence_keywords = ["证人证言", "书证", "物证", "鉴定意见", "监控视频"]
    evidences = [kw for kw in evidence_keywords if kw in text]
    procedure_nodes = [n for n in ["立案", "侦查", "移送审查起诉", "提起公诉"] if n in text]

    confidence = 0.92 if file_path.suffix.lower() in {".txt", ".md", ".json"} else 0.75
    return ExtractedMaterial(
        people=list(dict.fromkeys(names[:8])),
        facts=facts or ["待人工确认：未自动识别出明确事实时间节点"],
        evidences=evidences or ["待人工确认：未识别到标准证据类型"],
        procedure_nodes=procedure_nodes or ["待人工确认：程序节点信息不足"],
        ocr_confidence=confidence,
        raw_preview=text[:280],
    )


def recommend_laws(case_text: str, law_items: List[Dict]) -> List[Dict]:
    now = date.today().isoformat()
    suggestions: List[Dict] = []
    for item in law_items:
        matched = any(k in case_text for k in item.get("keywords", []))
        is_expired = bool(item.get("valid_to") and item["valid_to"] < now)
        if matched:
            suggestions.append(
                {
                    "title": item["title"],
                    "article": item["article"],
                    "source": item["source"],
                    "warning": "法条可能过时" if is_expired else "",
                }
            )
    return suggestions


def build_report_content(case: Dict, template: Dict, references: List[Dict], laws: List[Dict]) -> GenerationResult:
    extracted_blocks = [m["extracted"] for m in case.get("materials", [])]
    indictment = case.get("indictment", {}).get("extracted", {})

    all_people: List[str] = []
    all_facts: List[str] = []
    all_evidences: List[str] = []
    all_procedure_nodes: List[str] = []

    for block in extracted_blocks + ([indictment] if indictment else []):
        all_people.extend(block.get("people", []))
        all_facts.extend(block.get("facts", []))
        all_evidences.extend(block.get("evidences", []))
        all_procedure_nodes.extend(block.get("procedure_nodes", []))

    people = "、".join(list(dict.fromkeys(all_people))[:10]) or "待人工确认"
    facts = "；".join(list(dict.fromkeys(all_facts))[:5])
    evidences = "；".join(list(dict.fromkeys(all_evidences))[:8])
    procedures = "→".join(list(dict.fromkeys(all_procedure_nodes))[:5])

    style_hint = "；".join(r["content"][:40] for r in references[:2]) or "语言应严谨、客观、规范"

    sections = []
    for section_name in template.get("sections", []):
        if "事实" in section_name:
            content = f"经审查，涉及人员：{people}。主要事实：{facts}。"
        elif "证据" in section_name:
            content = f"证据目录与摘要：{evidences}。对证据三性需进一步人工审查。"
        elif "法律" in section_name or "适用" in section_name:
            content = "建议适用法律依据：" + "；".join(f"{l['title']}{l['article']}" for l in laws) if laws else "待人工确认：法律适用依据不足"
        elif "程序" in section_name:
            content = f"程序性节点：{procedures}。"
        else:
            content = f"本节依据模板生成。参考文风：{style_hint}。"
        sections.append({"heading": section_name, "content": content})

    traceability = [
        {"conclusion": "主要事实", "source": "案件证据材料", "locator": "上传文件预览/第1-3页"},
        {"conclusion": "法律适用", "source": "法律库", "locator": "法条推荐列表"},
    ]

    return GenerationResult(
        title=f"{case['title']}审查报告（草稿）",
        sections=sections,
        references_used=[r["title"] for r in references],
        law_suggestions=laws,
        traceability=traceability,
    )


def export_docx(report: Dict, target: Path) -> None:
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("请先安装 python-docx") from exc

    document = Document()
    document.add_heading(report["title"], level=1)
    for section in report["sections"]:
        document.add_heading(section["heading"], level=2)
        document.add_paragraph(section["content"])

    document.add_heading("引用依据清单", level=2)
    for law in report.get("law_suggestions", []):
        paragraph = f"{law['title']}{law['article']}（来源：{law['source']}）"
        if law.get("warning"):
            paragraph += f"；提示：{law['warning']}"
        document.add_paragraph(paragraph)

    document.add_heading("可追溯信息", level=2)
    for row in report.get("traceability", []):
        document.add_paragraph(f"{row['conclusion']} <- {row['source']} ({row['locator']})")

    document.save(target)
